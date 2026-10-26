#!/usr/bin/env python3
"""
In-season pipeline: UIL brackets -> season workbook -> public/data/{year}.json

  python3 scripts/season.py prefill --year 2026   # pull teams/scores into the workbook
  python3 scripts/season.py build   --year 2026   # workbook -> public/data/2026.json

The workbook (season/{year}/{year}_season.xlsx) is the source of truth. prefill only
adds games and fills blank teams/scores; it never touches the Chapter column, so
the chapter assignments typed into the workbook are always kept.

Bracket sources, in order:
  1. UIL bracket PDFs posted on uiltexas.org (parsed with uil_bracket.py)
  2. The MaxPreps bracket each UIL bracket page embeds (best effort; MaxPreps
     often refuses automated requests)
"""
import argparse, datetime, json, re, sys, tempfile, time, urllib.request
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from uil_bracket import parse as parse_pdf, norm, sim  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
CLASSES = ['1A D1', '1A D2', '2A D1', '2A D2', '3A D1', '3A D2',
           '4A D1', '4A D2', '5A D1', '5A D2', '6A D1', '6A D2']
ROUNDS = {1: 'Bi-District', 2: 'Area', 3: 'Regionals', 4: 'Quarterfinals', 5: 'Semifinals', 6: 'Championship'}
UIL = 'https://www.uiltexas.org'
UA = 'Mozilla/5.0 (compatible; texasplayoffs.com bracket sync)'
FIRST_ROW, LAST_ROW = 2, 801          # Games tab rows the template formats
COL = dict(week='A', cls='C', chapter='D', team1='E', score1='F', team2='G', score2='H', forfeit='I', notes='K')


def workbook_path(year):
    return ROOT / 'season' / str(year) / f'{year}_season.xlsx'


def fetch(url, binary=False):
    req = urllib.request.Request(url, headers={'User-Agent': UA})
    with urllib.request.urlopen(req, timeout=45) as r:
        data = r.read()
    return data if binary else data.decode('utf-8', 'replace')


# ── sources ──────────────────────────────────────────────────────────────

def uil_pdf_urls(year):
    """Bracket PDF links for this season found on UIL's football pages, newest per class."""
    pages = [f'{UIL}/football/playoff-brackets', f'{UIL}/football/archives'] + \
            [f'{UIL}/football/playoff-brackets/{c.replace(" ", "")}' for c in CLASSES]
    found = {}
    for p in pages:
        try:
            html = fetch(p)
        except Exception as e:
            print(f'  ! {p}: {e}')
            continue
        for href in re.findall(r'href="([^"]+\.pdf)"', html, re.I):
            m = re.search(rf'brackets/football/{year}_(\dA)-(D\d)[^/]*\.pdf$', href, re.I)
            if m:
                cls = f'{m.group(1).upper()} {m.group(2).upper()}'
                url = href if href.startswith('http') else UIL + href
                found.setdefault(cls, set()).add(url)
    # several revisions may be posted (…Bracket-1.pdf, …Bracket-2.pdf); take the highest
    rev = lambda u: [int(n) for n in re.findall(r'(\d+)\.pdf$', u)] or [0]
    return {c: max(urls, key=rev) for c, urls in found.items()}


# bracket slots that aren't a real team yet ("Winner G12", "W3", "TBD", ...)
PLACEHOLDER = re.compile(r'^((winner|loser)\s+(of\s+)?g(ame)?\s*\d+|[wr]\s*\d+|tbd|tba|bye|\?)$', re.I)


def real_teams(g):
    return all(g[k] and not PLACEHOLDER.match(g[k].strip()) for k in ('team1', 'team2'))


def games_from_pdf(url, cls):
    with tempfile.NamedTemporaryFile(suffix='.pdf') as f:
        f.write(fetch(url, binary=True)); f.flush()
        games, issues = parse_pdf(f.name)
    for i in issues:
        if 'state final' not in i:
            print(f'  ! {cls}: {i}')
    out = []
    for g in games:
        s1 = s2 = None
        if g['ws'] is not None:
            s1, s2 = (g['ws'], g['ls']) if g['winner'] == g['team1'] else (g['ls'], g['ws'])
        out.append(dict(week=g['round'], cls=cls, team1=g['team1'], team2=g['team2'], score1=s1, score2=s2))
    return out


def games_from_maxpreps(cls, year):
    from bs4 import BeautifulSoup
    page = fetch(f'{UIL}/football/playoff-brackets/{cls.replace(" ", "")}')
    m = re.search(r'https://www\.maxpreps\.com/tournament/view\.aspx[^"]+', page)
    if not m:
        return []
    soup = BeautifulSoup(fetch(m.group(0).replace('&amp;', '&')), 'html.parser')
    title = soup.title.get_text() if soup.title else ''
    if str(year) not in title:
        print(f'  ! {cls}: MaxPreps bracket is for another season ({title.strip()[:60]})')
        return []
    starts = sorted({a['title'][:10] for a in soup.select('abbr.start-date[title]')})
    if not starts:
        return []
    first = datetime.date.fromisoformat(starts[0])
    out = []
    for rnd in soup.select('div.round'):
        d = rnd.select_one('abbr.start-date[title]')
        if not d:
            continue
        week = min(6, 1 + (datetime.date.fromisoformat(d['title'][:10]) - first).days // 7)
        for mu in rnd.select('li.matchup-container'):
            teams = mu.select('li.team')
            if len(teams) != 2:
                continue
            names = [(t.select_one('.name').get_text(strip=True) if t.select_one('.name') else '') for t in teams]
            if not all(names):
                continue
            scores = []
            for t in teams:
                el = t.select_one('[class*="score"]')
                txt = el.get_text(strip=True) if el else ''
                scores.append(int(txt) if txt.isdigit() else None)
            out.append(dict(week=week, cls=cls, team1=names[0], team2=names[1], score1=scores[0], score2=scores[1]))
    return out


# ── workbook ─────────────────────────────────────────────────────────────

def open_workbook(year, create=False):
    from openpyxl import load_workbook
    path = workbook_path(year)
    if not path.exists():
        if not create:
            sys.exit(f'No workbook at {path}')
        import subprocess
        path.parent.mkdir(parents=True, exist_ok=True)
        subprocess.run([sys.executable, str(Path(__file__).parent / 'season_template.py'), str(path), str(year)], check=True)
    return load_workbook(path), path


def read_rows(ws):
    rows = []
    for r in range(FIRST_ROW, LAST_ROW + 1):
        v = {k: ws[f'{c}{r}'].value for k, c in COL.items()}
        v['row'] = r
        for k in ('chapter', 'team1', 'team2', 'cls', 'forfeit', 'notes'):
            v[k] = str(v[k]).strip() if v[k] not in (None, '') else None
        rows.append(v)
    return rows


def prefill(year):
    wb, path = open_workbook(year, create=True)
    ws = wb['Games']
    rows = read_rows(ws)
    fetched = {}
    pdfs = uil_pdf_urls(year)
    for cls in CLASSES:
        if cls in pdfs:
            try:
                fetched[cls] = games_from_pdf(pdfs[cls], cls)
                print(f'  {cls}: {len(fetched[cls])} games from {pdfs[cls].rsplit("/", 1)[-1]}')
                continue
            except Exception as e:
                print(f'  ! {cls}: PDF failed ({e})')
        try:
            fetched[cls] = games_from_maxpreps(cls, year)
            print(f'  {cls}: {len(fetched[cls])} games from MaxPreps')
        except Exception as e:
            print(f'  ! {cls}: no bracket available ({e})')
        time.sleep(8)  # be gentle with MaxPreps

    added = updated = 0
    conflicts = []
    for cls, games in fetched.items():
        for g in games:
            if not real_teams(g):
                continue
            if cls.startswith('1A') and g['week'] == 5:
                g['week'] = 6   # 1A's final is a week early; file all state finals under Championship
            match = None
            for r in rows:
                if r['week'] == g['week'] and r['cls'] == cls and r['team1'] and r['team2']:
                    straight = min(sim(r['team1'], g['team1']), sim(r['team2'], g['team2']))
                    crossed = min(sim(r['team1'], g['team2']), sim(r['team2'], g['team1']))
                    if max(straight, crossed) >= 0.85:
                        match = (r, straight >= crossed)
                        break
            if match is None:
                r = next((r for r in rows if not r['week'] and not r['team1'] and not r['team2'] and not r['cls']), None)
                if r is None:
                    sys.exit('Games tab is full')
                for k, v in (('week', g['week']), ('cls', cls), ('team1', g['team1']), ('team2', g['team2']),
                             ('score1', g['score1']), ('score2', g['score2'])):
                    ws[f'{COL[k]}{r["row"]}'] = v
                    r[k] = v
                added += 1
                continue
            r, straight = match
            s1, s2 = (g['score1'], g['score2']) if straight else (g['score2'], g['score1'])
            if s1 is None:
                continue
            if r['score1'] is None and r['score2'] is None:
                ws[f'{COL["score1"]}{r["row"]}'], ws[f'{COL["score2"]}{r["row"]}'] = s1, s2
                r['score1'], r['score2'] = s1, s2
                updated += 1
            elif (r['score1'], r['score2']) != (s1, s2):
                conflicts.append(f'row {r["row"]} {cls} wk{g["week"]}: workbook {r["team1"]} {r["score1"]}-{r["score2"]} {r["team2"]}, UIL {s1}-{s2}')
    if added or updated:
        wb.calculation.fullCalcOnLoad = True
        wb.save(path)
    print(f'prefill: {added} games added, {updated} scores filled')
    for c in conflicts:
        print('  score differs (left as is):', c)


def build(year, force=False):
    wb, path = open_workbook(year)
    ws = wb['Games']
    names = {}
    for r in range(2, 200):
        code, name = wb['Lists'][f'A{r}'].value, wb['Lists'][f'B{r}'].value
        if code:
            names[str(code).strip()] = name
    games = []
    for r in read_rows(ws):
        if not (r['week'] and r['cls'] and r['team1'] and r['team2']):
            continue
        week = int(r['week'])
        s1, s2 = r['score1'], r['score2']
        forfeit = (r['forfeit'] or '').upper() == 'Y'
        winner = r['team1'] if (s1 is not None and s2 is not None and s1 > s2) else \
                 r['team2'] if (s1 is not None and s2 is not None and s2 > s1) else None
        games.append(dict(year=year, classification=r['cls'], week=week, round=ROUNDS[week],
                          chapter=None if forfeit else r['chapter'], team1=r['team1'], score1=s1,
                          team2=r['team2'], score2=s2, winner=winner))
    games.sort(key=lambda g: (g['week'], CLASSES.index(g['classification']) if g['classification'] in CLASSES else 99))

    chapters = {}
    for g in games:
        if not g['chapter']:
            continue
        c = chapters.setdefault(g['chapter'], dict(code=g['chapter'], name=names.get(g['chapter']) or g['chapter'], total=0, weeks={}))
        w = c['weeks'].setdefault(str(g['week']), dict(total=0, classifications={}))
        c['total'] += 1; w['total'] += 1
        w['classifications'][g['classification']] = w['classifications'].get(g['classification'], 0) + 1
    chapter_list = sorted(chapters.values(), key=lambda c: -c['total'])

    played = sum(1 for g in games if g['score1'] is not None)
    assigned = sum(1 for g in games if g['chapter'])
    finals = sum(1 for g in games if g['week'] == 6 and g['score1'] is not None)
    out = dict(year=year,
               season=dict(status='complete' if finals >= 10 else 'in_progress',
                           updated=datetime.date.today().isoformat(),
                           games=len(games), played=played, assigned=assigned,
                           weeks={str(w): dict(games=sum(1 for g in games if g['week'] == w),
                                               assigned=sum(1 for g in games if g['week'] == w and g['chapter']))
                                  for w in range(1, 7)}),
               chapters=chapter_list, games=games)
    dest = ROOT / 'public' / 'data' / f'{year}.json'
    if dest.exists() and 'season' not in json.loads(dest.read_text()) and not force:
        sys.exit(f'{dest.name} was not made by this script (hand-curated season); refusing to overwrite. Use --force.')
    if dest.exists():
        # keep the old "updated" date when nothing else changed, so reruns don't create commits
        old = json.loads(dest.read_text())
        if old.get('season', {}).get('updated'):
            old_updated = old['season']['updated']
            old['season']['updated'] = out['season']['updated']
            if old == out:
                out['season']['updated'] = old_updated
    dest.write_text(json.dumps(out, indent=2) + '\n')
    idx_path = ROOT / 'public' / 'data' / 'index.json'
    idx = json.loads(idx_path.read_text())
    if year not in idx['years']:
        idx['years'] = sorted(idx['years'] + [year])
        idx_path.write_text(json.dumps(idx, indent=2) + '\n')
    print(f'build: {len(games)} games ({played} played, {assigned} with a chapter) -> {dest.relative_to(ROOT)}')


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('command', choices=['prefill', 'build'])
    today = datetime.date.today()
    # the playoffs run Nov-Dec; Jan-Jun still belongs to the previous season
    ap.add_argument('--year', type=int, default=today.year if today.month >= 7 else today.year - 1)
    ap.add_argument('--force', action='store_true', help='build: overwrite a hand-curated year file')
    a = ap.parse_args()
    prefill(a.year) if a.command == 'prefill' else build(a.year, a.force)
